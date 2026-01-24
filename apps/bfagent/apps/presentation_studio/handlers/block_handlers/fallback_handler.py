"""Fallback Block Handler - Handles unknown block types"""
from pptx.util import Inches, Pt
from .base_handler import BaseBlockHandler
import logging

logger = logging.getLogger(__name__)


class FallbackHandler(BaseBlockHandler):
    """Handler for unknown block types - text extraction fallback"""
    
    def render(self, content, styling, layout, block_type="unknown"):
        """Render fallback block with text extraction"""
        # Calculate safe height that fits within slide bounds
        desired_height = 3.5
        safe_height = self._calculate_safe_height(desired_height)
        
        box = self.slide.shapes.add_textbox(
            Inches(0.5), 
            Inches(self.current_y), 
            Inches(9), 
            Inches(safe_height)
        )
        frame = box.text_frame
        frame.word_wrap = True
        
        # Add block type indicator
        p = frame.paragraphs[0]
        p.text = f"[{block_type.upper().replace('_', ' ')}]"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.italic = True
        
        # Extract all text from content recursively
        texts = self._extract_all_text(content)
        
        # Limit text items based on available height
        max_items = int(safe_height * 5)  # ~5 items per inch
        for text in texts[:max_items]:
            p = frame.add_paragraph()
            p.text = f"• {text}"
            p.font.size = Pt(12)
            p.level = 1
        
        if len(texts) > max_items:
            p = frame.add_paragraph()
            p.text = f"... (+{len(texts) - max_items} more items)"
            p.font.size = Pt(10)
            p.font.italic = True
        
        logger.warning(f"Used fallback handler for block type: {block_type}")
        
        # Return next Y position with small gap
        return self.current_y + safe_height + 0.3
    
    def _extract_all_text(self, obj, texts=None):
        """Recursively extract all text strings from nested dict/list"""
        if texts is None:
            texts = []
        
        if isinstance(obj, dict):
            for key, value in obj.items():
                # Skip certain keys
                if key in ['styling', 'layout', 'animation', 'position', 'color', 'size']:
                    continue
                self._extract_all_text(value, texts)
        elif isinstance(obj, list):
            for item in obj:
                self._extract_all_text(item, texts)
        elif isinstance(obj, str) and obj.strip() and len(obj) > 2:
            # Only add meaningful strings
            if not obj.startswith('#') and not obj.startswith('http'):
                texts.append(obj.strip())
        
        return texts
